"""Reproduce bounded cobalt source comparisons; never publish or approve data.

The numeric transcriptions below were reviewed by the agent against PDF pages
or official HTML. Hash validation preserves that evidence; it is not independent
proof of a transcription, nor a human review. Page numbers are PDF, one-based.
"""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from . import archive, cobalt_collect

ROOT = Path(__file__).resolve().parents[1] / 'data/review/cobalt'


def main():
    sources = {e['id']: e['raw'] for e in json.loads((ROOT / 'external-sources.json').read_bytes())}
    for e in json.loads((ROOT / 'a0-probe.json').read_bytes()):
        if e.get('source') == 'usgs-cobalt':
            edition = '2025' if 'mcs2025' in e['url'] else '2026'
            sources['usgs-' + edition] = e['raw']
    for ref in sources.values():
        path = ROOT / ref['path']
        if hashlib.sha256(path.read_bytes()).hexdigest() != ref['sha256']:
            raise ValueError('Source hash changed: ' + str(path))
        if json.loads(path.with_suffix('.meta.json').read_bytes()) != ref:
            raise ValueError('Source sidecar differs: ' + str(path))

    observations = []

    def add(source, locator, year, entity, value, measure, method='agent_visual_review', **flags):
        observations.append(dict(source=source, locator=locator, year=year,
                                 entity=entity, value=value, unit='t Co', measure=measure,
                                 review_method=method, human_reviewed=False,
                                 publishable=False, **flags))

    # All 2024 country rows in the two USGS editions, preserving changed residual scope.
    editions = {
        '2025': dict(USA=300, AUS=3600, CAN=4500, COD=220000, CUB=3500,
                     IDN=28000, MDG=2600, NCL=1500, PNG=2800, PHL=3800,
                     RUS=8700, TUR=2700, OTHER=6200, WORLD=290000),
        '2026': dict(USA=200, AUS=4780, CAN=3350, CHN=2000, COD=226000,
                     CUB=3450, IDN=35000, MDG=3100, PNG=2630, PHL=3100,
                     RUS=8000, TUR=2200, OTHER=7780, WORLD=302000),
    }
    for edition, values in editions.items():
        for entity, value in values.items():
            add('usgs-' + edition, 'PDF p2, World Mine Production and Reserves, 2024 column',
                2024, entity, value, 'mine_contained_cobalt', source_estimated=True,
                scope='country' if entity not in ('OTHER', 'WORLD') else 'aggregate')
    bgs = dict(FIN=1008, RUS=8200, TUR=425, COD=200253, MDG=2535, MAR=1286,
               ZAF=611, ZMB=1410, ZWE=344, CAN=3351, CUB=3500, USA=300,
               BRA=0, CHN=1500, IDN=31570, PHL=3050, VNM=195, AUS=4776,
               NCL=1608, PNG=2625, WORLD=269000)
    for entity, value in bgs.items():
        add('bgs-2020-2024', 'PDF p25 / printed p15, Mine production of cobalt, 2024',
            2024, entity, value, 'metal_refined' if entity in ('MAR', 'ZAF') else 'mine_mixed_contained_or_recovered',
            source_estimated=entity in ('RUS', 'CUB', 'USA', 'CHN', 'IDN', 'PHL'),
            scope='aggregate' if entity == 'WORLD' else 'country',
            raw_zero_marker='dash' if entity == 'BRA' else None)
    for year, value in [(2024, 3351), (2023, 3260)]:
        add('nrcan-cobalt', 'Canadian production / text version, mine production', year,
            'CAN', value, 'mine_cobalt_in_concentrate', method='agent_html_review')
    add('bgs-2020-2024', 'PDF p25 / printed p15, Canada 2023', 2023, 'CAN', 5099,
        'mine_mixed_contained_or_recovered')
    add('usgs-2025', 'PDF p2, Canada 2023', 2023, 'CAN', 4220, 'mine_contained_cobalt')
    add('drc-ministry-2024', 'PDF p22, Table 32, total', 2024, 'COD', 198844.05,
        'production_sold_including_local_sales')
    for year, value in [(2024, 198777.21), (2023, 152798.86)]:
        add('drc-ministry-2024', 'PDF p26, Table 35, total cobalt metal', year,
            'COD', value, 'exports_contained_cobalt')
    add('drc-ministry-2023', 'PDF p17, Tables 15 and 16', 2023, 'COD', 139840.09,
        'exports_cobalt_tonnes_as_labeled')
    add('drc-ministry-index', 'Dashboard: Cobalt exporté (2024)', 2024, 'COD', 139840,
        'exports_dashboard_label', method='agent_html_review', disputed_year=True)
    add('glencore-2024', 'PDF p4, African Copper (KCC, Mutanda), cobalt 35.1 kt',
        2024, 'KCC+Mutanda', 35100, 'company_production_contained_in_concentrates_and_hydroxides')
    add('cmoc-2024-results', 'Official results release 2025-03-24, cobalt production',
        2024, 'CMOC', 114165, 'company_cobalt_production', method='agent_html_review')

    correlations = []
    # Verify every matching workbook row, not just the first match. Numeric or textual cells work.
    codes = {'260500', '282200', '810520'}
    for source in ('un-h6-h5', 'un-h6-h4', 'un-h5-h4'):
        wb = load_workbook(ROOT / sources[source]['path'], read_only=True, data_only=True)
        for sheet in wb:
            if 'corr' not in sheet.title.lower() and 'conv' not in sheet.title.lower():
                continue
            found = set()
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                vals = [str(v).strip() if v is not None else '' for v in row]
                targets = codes.intersection(vals)
                if not targets:
                    continue
                if len(targets) != 1:
                    raise ValueError('Cobalt mapping crosses target codes')
                code = next(iter(targets))
                nonempty = [v for v in vals if v]
                expected = [code, code] + (['1:1'] if 'corr' in sheet.title.lower() else [])
                if nonempty != expected or code in found:
                    raise ValueError('Non-identity or duplicate HS mapping: ' + str(nonempty))
                found.add(code)
                correlations.append(dict(source=source, sheet=sheet.title, row=row_number,
                                         code=code, cells=vals, mapping='1:1'))
            if found != codes:
                raise ValueError('Missing target HS mapping')
        wb.close()
    if len(correlations) != 18:
        raise ValueError('Unexpected number of target workbook mappings')

    groups = cobalt_collect.replay(ROOT / 'auth')
    def trade(year, code, reporter, partner, flow):
        matches = [r for key, rows in groups.items() for r in rows
                   if json.loads(key)['reporterCode'] != '' and
                   (r['year'], r['hs_code'], r['reporter'], r['partner'], r['flow']) ==
                   (year, code, reporter, partner, flow)]
        if len(matches) != 1:
            raise ValueError('Trade comparison is not unique')
        return matches[0]
    china = trade(2017, '260500', 'CHN', 'COD', 'M')
    drc = trade(2017, '260500', 'COD', 'CHN', 'X')
    calculations = {
        'usgs_2024_revision': {k: dict(old=editions['2025'][k], new=v,
            delta=v-editions['2025'][k], percent=100*(v/editions['2025'][k]-1))
            for k, v in editions['2026'].items() if k in editions['2025'] and k != 'OTHER'},
        'usgs_country_sum_minus_rounded_world': {edition: sum(v for k,v in values.items() if k != 'WORLD')-values['WORLD']
                                               for edition, values in editions.items()},
        'drc_2023_export_revision_percent': 100*(152798.86/139840.09-1),
        'canada_2024_usgs_2026_minus_nrcan_t': 3350-3351,
        'mirror_2017_260500': dict(china_import=china, drc_export=drc,
             import_over_export_weight=china['net_weight_kg']/drc['net_weight_kg'],
             import_over_export_value=china['value_usd']/drc['value_usd']),
        'china_2024_810520_from_drc': trade(2024, '810520', 'CHN', 'COD', 'M'),
    }
    output = dict(review_date='2026-09-15', publishable=False, human_reviewed=False,
                  scope='bounded 2017/2024 diagnostic and selected 2023/2024 source comparisons',
                  sources=sources, observations=observations, hs_correlations=correlations,
                  calculations=calculations)
    (ROOT / 'source-review.json').write_bytes(archive.encode(output))
    print(f'Verified {len(sources)} source hashes/sidecars, {len(correlations)} HS rows; reproduced {len(observations)} reviewed observations.')


if __name__ == '__main__':
    main()
